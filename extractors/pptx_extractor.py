from pptx import Presentation

def extract_pptx(file_path):
    presentation = Presentation(file_path)

    results = []

    for slide_number, slide in enumerate(
        presentation.slides,
        start=1
    ):

        for shape_number, shape in enumerate(
            slide.shapes,
            start=1
        ):

            if hasattr(shape, "text"):

                text = shape.text.strip()

                if text:
                    results.append({
                        "text": text,
                        "source": str(file_path),
                        "location": {
                            "type": "slide",
                            "slide": slide_number,
                            "shape": shape_number
                        },
                        "extraction_method": "native"
                    })

    return results
