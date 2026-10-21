# extractor.py

from pathlib import Path

from docx import Document
from pptx import Presentation
import fitz  # PyMuPDF
import pytesseract


def extract_docx(file_path):
    document = Document(file_path)

    results = []

    # Paragraphs
    for i, paragraph in enumerate(document.paragraphs):

        text = paragraph.text.strip()

        if text:
            results.append({
                "text": text,
                "source": str(file_path),
                "location": {
                    "type": "paragraph",
                    "index": i
                },
                "extraction_method": "native"
            })

    # Tables
    for table_index, table in enumerate(document.tables):

        for row_index, row in enumerate(table.rows):

            row_text = []

            for cell in row.cells:
                cell_text = cell.text.strip()

                if cell_text:
                    row_text.append(cell_text)

            if row_text:
                results.append({
                    "text": " | ".join(row_text),
                    "source": str(file_path),
                    "location": {
                        "type": "table",
                        "table": table_index,
                        "row": row_index
                    },
                    "extraction_method": "native"
                })

    return results
