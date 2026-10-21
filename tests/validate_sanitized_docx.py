import re
from pathlib import Path
from docx import Document

BASE_DIR = Path(__file__).resolve().parent
DATA_DIR = BASE_DIR / "data"

SANITIZED_FILE = (
    BASE_DIR
    / "output"
    / "Cadence_TPRM_Training_sanitized.docx"
)

MAPPING_FILE = DATA_DIR / "pseudonym_mapping.json"


def get_document_text(document):
    parts = []

    for paragraph in document.paragraphs:
        parts.append(paragraph.text)

    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)

    for section in document.sections:
        for paragraph in section.header.paragraphs:
            parts.append(paragraph.text)

        for paragraph in section.footer.paragraphs:
            parts.append(paragraph.text)

    return "\n".join(parts)


def main():
    if not SANITIZED_FILE.exists():
        raise FileNotFoundError(SANITIZED_FILE)

    document = Document(SANITIZED_FILE)
    text = get_document_text(document)

    print("Document opened successfully.")
    print(f"Characters inspected: {len(text)}")

    # Basic structural checks for common PII patterns.
    email_pattern = r"\b[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}\b"

    phone_pattern = (
        r"(?<!\w)(?:\+?\d[\d\s().-]{7,}\d)(?!\w)"
    )

    emails = re.findall(email_pattern, text)
    phones = re.findall(phone_pattern, text)

    print(f"Email-like values remaining: {len(emails)}")
    print(f"Phone-like values remaining: {len(phones)}")

    if emails:
        print("Email examples:", emails[:5])

    if phones:
        print("Phone examples:", phones[:5])

    # Check whether original mapped values remain.
    ENTITIES_FILE = DATA_DIR / "pseudonymized_entities.json"

    if ENTITIES_FILE.exists():
        import json

        with open(ENTITIES_FILE, "r", encoding="utf-8") as file:
            entities = json.load(file)

        remaining = []

        for entity in entities:
            for mention in entity.get("mentions", []):
                original = mention.get("text", "").strip()

                if len(original) < 4:
                    continue

                if original.casefold() in text.casefold():
                    remaining.append({
                        "text": original,
                        "type": entity.get("type"),
                        "pseudonym": entity.get("pseudonym"),
                    })

        # Remove duplicate reports.
        unique_remaining = {
            (item["text"].casefold(), item["type"]): item
            for item in remaining
        }

        remaining = list(unique_remaining.values())

        print(
            "Detected original mentions still present: "
            f"{len(remaining)}"
        )

        for item in remaining[:20]:
            print(
                f"Text: {item['text']!r} | "
                f"Type: {item['type']} | "
                f"Expected token: {item['pseudonym']}"
            )

    print("\nValidation checks completed.")
    print("Review any reported matches before declaring the document safe.")


if __name__ == "__main__":
    main()
